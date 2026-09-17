#!/usr/bin/env python3

import sys
import os
import json
import time
import argparse
import subprocess
import html
import shutil
from pathlib import Path
import hashlib
from html.parser import HTMLParser

from datetime import datetime

import requests

try:
    from tqdm import tqdm
except ImportError:  # pragma: no cover - optional dependency fallback
    def tqdm(iterable=None, **kwargs):
        return iterable if iterable is not None else []

from PIL import Image, ImageOps, ExifTags, ImageFile, ImageFilter, ImageStat

try:
    from ultralytics import YOLO
except Exception:  # pragma: no cover - optional dependency
    YOLO = None

try:
    import torch
except Exception:  # pragma: no cover - optional dependency
    torch = None

try:
    import torchvision.models as models
    from torchvision import transforms
except Exception:  # pragma: no cover - optional dependency
    models = None
    transforms = None

try:
    from transformers import AutoProcessor, AutoModel
except Exception:  # pragma: no cover - optional dependency
    AutoProcessor = None
    AutoModel = None

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.drawing.image import Image as XLImage


ImageFile.LOAD_TRUNCATED_IMAGES = True

# ======================================================
# Configuration
# ======================================================

INAT_API_URL = (
    "https://api.inaturalist.org/v2"
)

TOKEN_FILE = (
    "token.txt"
)

CHECKPOINT_FILE = (
    "inat_checkpoint.json"
)

LOCAL_CACHE_FILE = (
    "local_image_cache.json"
)

PHOTO_DETECTION_CACHE_FILE = (
    "photo_detection_cache.json"
)


SUPPORTED_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png"
}


USER_AGENT = (
    "iNaturalist-local-image-comparator/3.0"
)


PAGE_SIZE = 200
DEBUG_MODE = False

API_DELAY_SECONDS = 3


THUMB_SIZE = (
    120,
    120
)



# ======================================================
# Token
# ======================================================


def load_api_token():

    if not os.path.exists(
        TOKEN_FILE
    ):

        raise RuntimeError(
            "Missing token.txt\n"
            "Create token.txt and paste "
            "your iNaturalist API token."
        )


    with open(
        TOKEN_FILE,
        "r",
        encoding="utf-8"
    ) as f:

        token = f.read().strip()


    if not token:

        raise RuntimeError(
            "token.txt is empty"
        )


    return token



# ======================================================
# API Client
# ======================================================


class INaturalistClient:


    def __init__(
            self,
            token):


        self.session = requests.Session()


        self.session.headers.update(
            {
                "Authorization":
                    f"Bearer {token}",

                "User-Agent":
                    USER_AGENT
            }
        )



    def get(
            self,
            endpoint,
            params=None,
            retries=8):


        url = (
            INAT_API_URL
            +
            endpoint
        )


        attempt = 0


        while True:


            try:


                response = self.session.get(
                    url,
                    params=params,
                    timeout=60
                )

                print("Response URL: ", response.url)

                if response.status_code in (
                    403,
                    429
                ):


                    wait = (
                        30
                        *
                        (attempt + 1)
                    )


                    print(
                        f"API {response.status_code}. "
                        f"Waiting {wait}s"
                    )


                    time.sleep(
                        wait
                    )


                    attempt += 1


                    if attempt >= retries:

                        raise RuntimeError(
                            "API retry limit reached"
                        )


                    continue



                response.raise_for_status()


                time.sleep(
                    API_DELAY_SECONDS
                )


                return response.json()



            except requests.exceptions.RequestException as e:


                attempt += 1


                if attempt >= retries:

                    raise


                wait = (
                    30
                    *
                    attempt
                )


                print(
                    f"Request failed: {e}"
                )

                print(
                    f"Retrying after {wait}s"
                )


                time.sleep(
                    wait
                )



# ======================================================
# JSON helpers
# ======================================================


def save_json(
        filename,
        data):


    with open(
        filename,
        "w",
        encoding="utf-8"
    ) as f:

        json.dump(
            data,
            f,
            ensure_ascii=False,
            indent=2
        )



def load_json(
        filename):


    with open(
        filename,
        encoding="utf-8"
    ) as f:

        return json.load(f)


class _DetectionReportParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []

    def handle_starttag(self, tag, attrs):
        if tag.lower() != "tr":
            return
        attributes = dict(attrs)
        if "data-path" not in attributes:
            return
        self.rows.append({
            "path": attributes.get("data-path", ""),
            "matched": attributes.get("data-matched") == "true",
            "human": attributes.get("data-human") == "true",
            "landscape": attributes.get("data-landscape") == "true",
        })


def import_detection_cache_from_report(report_file, detection_cache_file):
    parser = _DetectionReportParser()
    with open(report_file, "r", encoding="utf-8") as f:
        parser.feed(f.read())

    detection_cache = {}
    imported = 0
    skipped_matched = 0
    skipped_missing = 0
    for row in parser.rows:
        if row["matched"]:
            skipped_matched += 1
            continue

        path = row["path"]
        if not path or not os.path.isfile(path):
            skipped_missing += 1
            continue

        detection_cache[_normalized_image_path(path)] = {
            "signature": get_file_signature(path),
            "flags": {
                "human": row["human"],
                "landscape": row["landscape"],
            },
        }
        imported += 1

    save_json(detection_cache_file, detection_cache)
    print("Imported photo detection results:", imported)
    print("Skipped matched report rows:", skipped_matched)
    if skipped_missing:
        print("Skipped report rows with missing files:", skipped_missing)
    print("Saved photo detection results to", detection_cache_file)



# ======================================================
# Filename helpers
# ======================================================


def normalize_filename(
        filename):

    if not filename:
        return None

    name = os.path.splitext(
        os.path.basename(filename)
    )[0].lower()

    for ch in ["_", "-", ",", ".", ";", ":"]:
        name = name.replace(ch, " ")
    return " ".join(name.split())



def timestamp_minute(
        dt):


    if not dt:

        return None


    return dt.strftime(
        "%Y-%m-%d %H:%M"
    )

import re

def parse_date_from_path(path):
    if not path:
        return None
    m = re.search(r"(\d{4})[-_.](\d{1,2})[-_.](\d{1,2})", path)
    if m:
        try:
            return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
        except Exception:
            pass
    m = re.search(r"(?:^|[^\d])(\d{1,2})[.](\d{1,2})[.](\d{2,4})(?:$|[^\d])", path)
    if m:
        d, m_num, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        if 1 <= m_num <= 12 and 1 <= d <= 31:
            return f"{y:04d}-{m_num:02d}-{d:02d}"
    return None

def get_candidate_date(candidate):
    ts = candidate.get("timestamp")
    if ts and len(ts) >= 10:
        return ts[:10]
    return parse_date_from_path(candidate.get("path"))

from PIL import Image

def get_thumbnail_filename(image_path):
    path_obj = Path(image_path).resolve()
    path_hash = hashlib.sha256(str(path_obj).encode("utf-8")).hexdigest()[:12]
    safe_stem = "".join(c for c in path_obj.stem if c.isalnum() or c in ("-", "_")).rstrip()
    if not safe_stem:
        safe_stem = "thumb"
    return f"{safe_stem}_{path_hash}.thumb-v2.jpg"


def create_thumbnail_file(
        source,
        output,
        size=150):

    if os.path.exists(output) and os.path.getsize(output) > 0:
        return True

    try:

        img = Image.open(source)
        img = ImageOps.exif_transpose(img)
        img.load()

        img.thumbnail(
            (size, size)
        )

        if img.mode in ("RGBA", "LA", "P"):

            background = Image.new(
                "RGB",
                img.size,
                (255, 255, 255)
            )

            if img.mode == "P":
                img = img.convert("RGBA")

            background.paste(
                img,
                mask=img.split()[-1]
            )

            img = background

        elif img.mode != "RGB":

            img = img.convert("RGB")

        img.save(
            output,
            "JPEG",
            quality=65
        )

        return True

    except Exception as e:

        print(
            f"Thumbnail failed: {source} ({e})"
        )

        return False

def esc(value):

    if value is None:
        return ""

    return html.escape(
        str(value)
    )


def sort_items(items, sort_by="date", sort_order="asc", group_by_dir=True, item_type="unmatched"):
    def sort_key(item):
        if item_type == "unmatched":
            path = item.get("path", "")
            dir_path = str(Path(path).parent).lower() if path else ""
            fname = (item.get("original_filename") or item.get("filename") or "").lower()
            ts = item.get("timestamp") or ""
        else:
            local_imgs = item.get("local_images", [])
            path = local_imgs[0] if local_imgs else ""
            dir_path = str(Path(path).parent).lower() if path else ""
            fname = (item.get("inat_original_filename") or "").lower()
            ts = item.get("timestamp") or ""

        dir_k = dir_path if group_by_dir else ""
        if sort_by == "filename":
            val_k = fname
        elif sort_by == "dir":
            val_k = dir_path
        else:
            val_k = ts if ts else ("0000-00-00" if sort_order == "desc" else "9999-99-99")

        return (dir_k, val_k, fname)

    reverse = (sort_order == "desc")
    return sorted(items, key=sort_key, reverse=reverse)


def create_html_report(
        results,
        local_missing,
        output_dir,
        sort_by="date",
        sort_order="asc",
        group_by_dir=True):

    output_dir = Path(output_dir)
    thumb_dir = output_dir / "thumbnails"
    local_thumb_dir = thumb_dir / "local"

    for d in [output_dir, local_thumb_dir]:
        d.mkdir(parents=True, exist_ok=True)

    html_file = output_dir / "report.html"

    matched_results = [r for r in results if r.get("status") == "MATCHED"]
    matched_results = sort_items(matched_results, sort_by=sort_by, sort_order=sort_order, group_by_dir=group_by_dir, item_type="matched")
    local_missing = sort_items(local_missing, sort_by=sort_by, sort_order=sort_order, group_by_dir=group_by_dir, item_type="unmatched")

    matched_count = len(matched_results)
    unmatched_count = len(local_missing)
    total_count = matched_count + unmatched_count

    thumbnail_paths = []
    seen_thumbnail_paths = set()
    for image in local_missing:
        path = image.get("path")
        normalized_path = _normalized_image_path(path)
        if normalized_path and normalized_path not in seen_thumbnail_paths:
            seen_thumbnail_paths.add(normalized_path)
            thumbnail_paths.append(path)
    for result in matched_results:
        for path in result.get("local_images", []):
            normalized_path = _normalized_image_path(path)
            if normalized_path and normalized_path not in seen_thumbnail_paths:
                seen_thumbnail_paths.add(normalized_path)
                thumbnail_paths.append(path)

    for path in tqdm(
            thumbnail_paths,
            desc="Creating report thumbnails",
            unit="thumbnail"):
        create_thumbnail_file(
            path,
            local_thumb_dir / get_thumbnail_filename(path)
        )

    camera_types = set()
    camera_names = set()
    for image in local_missing:
        if image.get("camera_type"):
            camera_types.add(image["camera_type"])
        if image.get("camera_name"):
            camera_names.add(image["camera_name"])
    for result in matched_results:
        for image in result.get("local_image_data", []):
            if image.get("camera_type"):
                camera_types.add(image["camera_type"])
            if image.get("camera_name"):
                camera_names.add(image["camera_name"])

    badge_css = '''
    .badge-human {{
        display: inline-block;
        padding: 4px 8px;
        font-size: 12px;
        font-weight: bold;
        color: #0b3d91;
        background-color: #d9ecff;
        border: 1px solid #b8d8ff;
        border-radius: 4px;
        margin-left: 8px;
        margin-top: 4px;
    }}
    .badge-landscape {{
        display: inline-block;
        padding: 4px 8px;
        font-size: 12px;
        font-weight: bold;
        color: #1f4d2e;
        background-color: #dff3d8;
        border: 1px solid #bfe4b7;
        border-radius: 4px;
        margin-left: 8px;
        margin-top: 4px;
    }}
    .badge-non-live {{
        display: inline-block;
        padding: 4px 8px;
        font-size: 12px;
        font-weight: bold;
        color: #856404;
        background-color: #ffeeba;
        border: 1px solid #ffe8a1;
        border-radius: 4px;
        margin-left: 8px;
        margin-top: 4px;
    }}
    '''

    camera_type_options = "".join(
        f'<option value="{esc(value)}">{esc(value)}</option>'
        for value in sorted(camera_types, key=str.casefold)
    )
    camera_name_options = "".join(
        f'<option value="{esc(value)}">{esc(value)}</option>'
        for value in sorted(camera_names, key=str.casefold)
    )

    sort_by_date_sel = 'selected' if sort_by == 'date' else ''
    sort_by_fn_sel = 'selected' if sort_by == 'filename' else ''
    sort_by_dir_sel = 'selected' if sort_by == 'dir' else ''

    sort_asc_sel = 'selected' if sort_order == 'asc' else ''
    sort_desc_sel = 'selected' if sort_order == 'desc' else ''

    group_by_dir_chk = 'checked' if group_by_dir else ''

    with open(html_file, "w", encoding="utf-8") as f:
        f.write(f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>iNaturalist comparison</title>
<style>
{badge_css}
body {{
    font-family: system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
    margin: 20px;
    background: #f8f9fa;
    color: #333;
}}
h1 {{
    color: #2c3e50;
    margin-bottom: 20px;
}}
.toolbar {{
    background: #fff;
    padding: 15px;
    border-radius: 8px;
    box-shadow: 0 2px 4px rgba(0,0,0,0.05);
    margin-bottom: 20px;
    display: flex;
    flex-wrap: wrap;
    gap: 15px;
    align-items: center;
}}
.toolbar input[type="text"] {{
    padding: 8px 12px;
    font-size: 15px;
    border: 1px solid #ccc;
    border-radius: 6px;
    width: 260px;
}}
.toolbar select {{
    padding: 8px 12px;
    font-size: 14px;
    border: 1px solid #ccc;
    border-radius: 6px;
    background: #fff;
    cursor: pointer;
}}
.btn {{
    padding: 8px 14px;
    font-size: 14px;
    border: none;
    border-radius: 6px;
    cursor: pointer;
    background: #007bff;
    color: white;
    font-weight: 500;
    transition: background 0.2s;
}}
.btn:hover {{
    background: #0056b3;
}}
.btn-success {{
    background: #28a745;
}}
.btn-success:hover {{
    background: #218838;
}}
.btn-secondary {{
    background: #6c757d;
}}
.btn-secondary:hover {{
    background: #5a6268;
}}
.btn-warning {{
    background: #ffc107;
    color: #212529;
}}
.btn-warning:hover {{
    background: #e0a800;
}}
.tabs {{
    display: flex;
    gap: 10px;
    margin-bottom: 15px;
    border-bottom: 2px solid #dee2e6;
    padding-bottom: 10px;
}}
.tab-btn {{
    padding: 10px 18px;
    font-size: 15px;
    border: none;
    border-radius: 6px 6px 0 0;
    background: #e9ecef;
    color: #495057;
    cursor: pointer;
    font-weight: bold;
}}
.tab-btn.active {{
    background: #28a745;
    color: white;
}}
table {{
    border-collapse: collapse;
    width: 100%;
    background: white;
    box-shadow: 0 2px 4px rgba(0,0,0,0.05);
    border-radius: 8px;
    overflow: hidden;
}}
th {{
    background: #f1f3f5;
    color: #495057;
    text-align: left;
    padding: 10px;
    border-bottom: 2px solid #dee2e6;
}}
td {{
    border-bottom: 1px solid #e9ecef;
    padding: 10px;
    vertical-align: middle;
}}
img {{
    max-width: 140px;
    max-height: 140px;
    border-radius: 4px;
    border: 1px solid #ddd;
    object-fit: cover;
}}
.pagination {{
    margin-top: 15px;
    display: flex;
    align-items: center;
    gap: 15px;
}}
.non-live-row {{
    background-color: #fff3cd !important;
    opacity: 0.7;
}}
.badge-non-live {{
    display: inline-block;
    padding: 4px 8px;
    font-size: 12px;
    font-weight: bold;
    color: #856404;
    background-color: #ffeeba;
    border: 1px solid #ffe8a1;
    border-radius: 4px;
    margin-left: 8px;
}}
</style>

<script>
let currentPage = 0;
const pageSize = 100;
let currentTab = "unmatched";

let nonLivePaths = new Set();
let nonLiveSignatures = new Set();

function initNonLiveStore() {{
    try {{
        let saved = localStorage.getItem("inat_non_live_photos");
        if (saved) {{
            let data = JSON.parse(saved);
            if (data.paths) data.paths.forEach(p => nonLivePaths.add(p));
            if (data.signatures) data.signatures.forEach(s => nonLiveSignatures.add(s));
        }}
    }} catch(e) {{
        console.error("Failed to load non-live photos from localStorage", e);
    }}
}}

function saveNonLiveStore() {{
    let payload = {{
        paths: Array.from(nonLivePaths),
        signatures: Array.from(nonLiveSignatures)
    }};
    localStorage.setItem("inat_non_live_photos", JSON.stringify(payload));
}}

function toggleNonLive(btn, path, filename, signature) {{
    let row = btn ? btn.closest("tr") : null;
    if (!path && row && row.dataset.path) path = row.dataset.path;
    if (!filename && row && row.dataset.filename) filename = row.dataset.filename;
    if (!signature && row && row.dataset.signature) signature = row.dataset.signature;

    let isNonLive = (path && nonLivePaths.has(path)) ||
                    (signature && nonLiveSignatures.has(signature));

    if (isNonLive) {{
        if (path) nonLivePaths.delete(path);
        if (signature) nonLiveSignatures.delete(signature);
        if (row) {{
            row.classList.remove("non-live-row");
            let badge = row.querySelector(".badge-non-live");
            if (badge) badge.style.display = "none";
        }}
        if (btn) {{
            btn.innerText = "🚫 Mark Not Live";
            btn.className = "btn btn-warning btn-toggle-nonlive";
        }}
    }} else {{
        if (path) nonLivePaths.add(path);
        if (signature) nonLiveSignatures.add(signature);
        if (row) {{
            row.classList.add("non-live-row");
            let badge = row.querySelector(".badge-non-live");
            if (badge) badge.style.display = "inline-block";
        }}
        if (btn) {{
            btn.innerText = "🌱 Mark as Live";
            btn.className = "btn btn-secondary btn-toggle-nonlive";
        }}
    }}
    saveNonLiveStore();
    filterTable();
}}

function exportNonLiveJSON() {{
    let payload = {{
        paths: Array.from(nonLivePaths),
        signatures: Array.from(nonLiveSignatures)
    }};
    let dataStr = "data:text/json;charset=utf-8," + encodeURIComponent(JSON.stringify(payload, null, 2));
    let downloadAnchor = document.createElement('a');
    downloadAnchor.setAttribute("href", dataStr);
    downloadAnchor.setAttribute("download", "non_live_photos.json");
    document.body.appendChild(downloadAnchor);
    downloadAnchor.click();
    downloadAnchor.remove();
}}

function importNonLiveJSON(event) {{
    let file = event.target.files[0];
    if (!file) return;
    let reader = new FileReader();
    reader.onload = function(e) {{
        try {{
            let data = JSON.parse(e.target.result);
            if (Array.isArray(data)) {{
                data.forEach(p => nonLivePaths.add(p));
            }} else if (typeof data === "object") {{
                if (data.paths) data.paths.forEach(p => nonLivePaths.add(p));
                if (data.signatures) data.signatures.forEach(s => nonLiveSignatures.add(s));
            }}
            saveNonLiveStore();
            applySavedNonLiveToRows();
            filterTable();
            alert("Loaded non-live creature choices successfully!");
        }} catch(err) {{
            alert("Error parsing JSON file: " + err.message);
        }}
    }};
    reader.readAsText(file);
}}

function applySavedNonLiveToRows() {{
    let rows = document.querySelectorAll("#results tbody tr");
    rows.forEach(function(row) {{
        let path = row.dataset.path;
        let sig = row.dataset.signature;
        let btn = row.querySelector(".btn-toggle-nonlive");
        let badge = row.querySelector(".badge-non-live");

        if (btn && (nonLivePaths.has(path) || (sig && nonLiveSignatures.has(sig)))) {{
            row.classList.add("non-live-row");
            btn.innerText = "🌱 Mark as Live";
            btn.className = "btn btn-secondary btn-toggle-nonlive";
            if (badge) badge.style.display = "inline-block";
        }}
    }});
}}

function switchTab(tab) {{
    currentTab = tab;
    document.querySelectorAll(".tab-btn").forEach(b => b.classList.remove("active"));
    let activeBtn = document.getElementById("tab-" + tab);
    if (activeBtn) activeBtn.classList.add("active");
    filterTable();
}}

function onSortChange() {{
    sortTable();
    filterTable();
}}

function sortTable() {{
    let tbody = document.querySelector("#results tbody");
    if (!tbody) return;
    let rows = Array.from(tbody.querySelectorAll("tr"));

    let sortBy = document.getElementById("sort-by").value;
    let sortOrder = document.getElementById("sort-order").value;
    let groupByDir = document.getElementById("group-by-dir").checked;

    rows.sort(function(a, b) {{
        let dirA = (a.dataset.dir || "").toLowerCase();
        let dirB = (b.dataset.dir || "").toLowerCase();

        let fileA = (a.dataset.filename || "").toLowerCase();
        let fileB = (b.dataset.filename || "").toLowerCase();

        let dateA = a.dataset.timestamp || "";
        let dateB = b.dataset.timestamp || "";

        if (groupByDir) {{
            let dirCmp = dirA.localeCompare(dirB);
            if (dirCmp !== 0) return dirCmp;
        }}

        let cmp = 0;
        if (sortBy === "filename") {{
            cmp = fileA.localeCompare(fileB);
        }} else if (sortBy === "date") {{
            if (!dateA && !dateB) cmp = 0;
            else if (!dateA) return 1;
            else if (!dateB) return -1;
            else cmp = dateA.localeCompare(dateB);
        }} else if (sortBy === "dir") {{
            cmp = dirA.localeCompare(dirB);
        }}

        if (sortOrder === "desc") {{
            cmp = -cmp;
        }}

        if (cmp === 0) {{
            cmp = fileA.localeCompare(fileB);
        }}
        return cmp;
    }});

    let fragment = document.createDocumentFragment();
    rows.forEach(function(row) {{
        fragment.appendChild(row);
    }});
    tbody.appendChild(fragment);
}}

function filterTable() {{
    let text = document.getElementById("search").value.toLowerCase();
    let hideNonLive = document.getElementById("hide-non-live").checked;
    let hideCurrentYear = document.getElementById("hide-current-year").checked;
    let currentYear = String(new Date().getFullYear());
    let flagFilter = document.getElementById("flag-filter").value;
    let selectedCameraTypes = Array.from(document.getElementById("camera-type-filter").selectedOptions).map(option => option.value);
    let selectedCameraNames = Array.from(document.getElementById("camera-name-filter").selectedOptions).map(option => option.value);
    let rows = document.querySelectorAll("#results tbody tr");

    rows.forEach(function(row) {{
        let isMatched = row.dataset.matched === "true";
        let isNonLive = row.classList.contains("non-live-row");
        let isHuman = row.dataset.human === "true";
        let isLandscape = row.dataset.landscape === "true";
        let isFlagged = isHuman || isLandscape;
        let isCurrentYear = (row.dataset.timestamp || "").slice(0, 4) === currentYear;

        let matchesTab = false;
        if (currentTab === "all") {{
            matchesTab = true;
        }} else if (currentTab === "unmatched" && !isMatched) {{
            matchesTab = true;
        }} else if (currentTab === "matched" && isMatched) {{
            matchesTab = true;
        }}

        let matchesSearch = !text || row.innerText.toLowerCase().includes(text);
        let matchesNonLiveFilter = !hideNonLive || !isNonLive;
        let matchesCurrentYearFilter = !hideCurrentYear || !isCurrentYear;
        let matchesFlagFilter = true;
        if (flagFilter === "flagged") matchesFlagFilter = isFlagged;
        else if (flagFilter === "human") matchesFlagFilter = isHuman;
        else if (flagFilter === "landscape") matchesFlagFilter = isLandscape;
        else if (flagFilter === "unflagged") matchesFlagFilter = !isFlagged;

        let rowCameraTypes = JSON.parse(row.dataset.cameraTypes || "[]");
        let rowCameraNames = JSON.parse(row.dataset.cameraNames || "[]");
        let matchesCameraType = !selectedCameraTypes.length || selectedCameraTypes.some(value => rowCameraTypes.includes(value));
        let matchesCameraName = !selectedCameraNames.length || selectedCameraNames.some(value => rowCameraNames.includes(value));

        row.dataset.hidden = !(matchesTab && matchesSearch && matchesNonLiveFilter && matchesCurrentYearFilter && matchesFlagFilter && matchesCameraType && matchesCameraName);
    }});

    showPage(currentPage);
}}

function showPage(page) {{
    currentPage = page;
    let rows = document.querySelectorAll("#results tbody tr");
    let visibleRows = [];

    rows.forEach(function(row) {{
        if (row.dataset.hidden !== "true") {{
            visibleRows.push(row);
        }}
        row.style.display = "none";
    }});

    let start = page * pageSize;
    let end = start + pageSize;

    visibleRows.slice(start, end).forEach(function(row) {{
        row.style.display = "";
    }});

    document.getElementById("page_info").innerText =
        "Page " + (page + 1) + " / " + Math.max(1, Math.ceil(visibleRows.length / pageSize)) +
        " (" + visibleRows.length + " visible)";
}}

function nextPage() {{
    let rows = document.querySelectorAll("#results tbody tr");
    let count = 0;
    rows.forEach(row => {{ if (row.dataset.hidden !== "true") count++; }});
    if ((currentPage + 1) * pageSize < count) showPage(currentPage + 1);
}}

function previousPage() {{
    if (currentPage > 0) showPage(currentPage - 1);
}}

window.onload = function() {{
    initNonLiveStore();
    applySavedNonLiveToRows();
    switchTab("unmatched");
}};
</script>
</head>

<body>

<h1>iNaturalist Local Comparison Report</h1>

<div class="toolbar">
    <input id="search" oninput="filterTable()" placeholder="Search filename, ID, date, path...">

    <label style="font-weight: 500; display: flex; align-items: center; gap: 5px;">
        Camera type:
        <select id="camera-type-filter" multiple size="3" onchange="filterTable()">
            {camera_type_options}
        </select>
    </label>

    <label style="font-weight: 500; display: flex; align-items: center; gap: 5px;">
        Camera name:
        <select id="camera-name-filter" multiple size="3" onchange="filterTable()">
            {camera_name_options}
        </select>
    </label>
    
    <label style="font-weight: 500; display: flex; align-items: center; gap: 5px;">
        Sort by:
        <select id="sort-by" onchange="onSortChange()">
            <option value="date" {sort_by_date_sel}>Date / Time</option>
            <option value="filename" {sort_by_fn_sel}>Filename</option>
            <option value="dir" {sort_by_dir_sel}>Directory Path</option>
        </select>
    </label>

    <label style="font-weight: 500; display: flex; align-items: center; gap: 5px;">
        Order:
        <select id="sort-order" onchange="onSortChange()">
            <option value="asc" {sort_asc_sel}>Ascending ⬆</option>
            <option value="desc" {sort_desc_sel}>Descending ⬇</option>
        </select>
    </label>

    <label style="font-weight: 500; cursor: pointer; display: flex; align-items: center; gap: 5px;">
        <input type="checkbox" id="group-by-dir" onchange="onSortChange()" {group_by_dir_chk}> Group by Directory
    </label>

    <label style="font-weight: 500; cursor: pointer; display: flex; align-items: center; gap: 5px;">
        <input type="checkbox" id="hide-non-live" onchange="filterTable()" checked> Hide non-live marked photos
    </label>

    <label style="font-weight: 500; cursor: pointer; display: flex; align-items: center; gap: 5px;">
        <input type="checkbox" id="hide-current-year" onchange="filterTable()"> Hide current year photos
    </label>

    <label style="font-weight: 500; display: flex; align-items: center; gap: 5px;">
        Auto-flag filter:
        <select id="flag-filter" onchange="filterTable()">
            <option value="all">All photos</option>
            <option value="human">Human only</option>
            <option value="landscape">Landscape only</option>
            <option value="unflagged">Unflagged only</option>
        </select>
    </label>
    <button class="btn btn-success" onclick="exportNonLiveJSON()">💾 Save Non-Live List (non_live_photos.json)</button>
    <input type="file" id="import-json" accept=".json" onchange="importNonLiveJSON(event)" style="display:none;">
    <button class="btn btn-secondary" onclick="document.getElementById('import-json').click()">📁 Load non_live_photos.json</button>
</div>

<div class="tabs">
    <button class="tab-btn active" id="tab-unmatched" onclick="switchTab('unmatched')">Unmatched Local Photos ({unmatched_count})</button>
    <button class="tab-btn" id="tab-matched" onclick="switchTab('matched')">Matched Observations ({matched_count})</button>
    <button class="tab-btn" id="tab-all" onclick="switchTab('all')">All Photos ({total_count})</button>
</div>

<div class="pagination">
    <button class="btn btn-secondary" onclick="previousPage()">Previous</button>
    <span id="page_info"></span>
    <button class="btn btn-secondary" onclick="nextPage()">Next</button>
</div>

<br>

<table id="results">
<thead>
<tr>
    <th>Status / Filename</th>
    <th>Thumbnail</th>
    <th>Path / iNat Details</th>
    <th>Timestamp / ID</th>
    <th>Actions</th>
</tr>
</thead>
<tbody>
""")

        # 1. Unmatched Local Photos
        for image in local_missing:
            thumb_name = get_thumbnail_filename(image["path"])

            img_path = esc(image.get("path", ""))
            dir_path = esc(str(Path(image.get("path", "")).parent))
            orig_name = esc(image.get("original_filename", ""))
            norm_filename = esc(image.get("filename", ""))
            timestamp = esc(image.get("timestamp", ""))
            camera_types_json = esc(json.dumps([image["camera_type"]] if image.get("camera_type") else []))
            camera_names_json = esc(json.dumps([image["camera_name"]] if image.get("camera_name") else []))
            human_badge = '<span class="badge-human">Includes Humans</span>' if image.get("human") else ''
            landscape_badge = '<span class="badge-landscape">Landscape / no clear subject</span>' if image.get("landscape") else ''

            is_human = "true" if image.get("human") else "false"
            is_landscape = "true" if image.get("landscape") else "false"
            badges_html = " ".join(part for part in [human_badge, landscape_badge] if part)

            f.write(f'<tr data-matched="false" data-path="{img_path}" data-dir="{dir_path}" data-filename="{orig_name}" data-timestamp="{timestamp}" data-camera-types="{camera_types_json}" data-camera-names="{camera_names_json}" data-human="{is_human}" data-landscape="{is_landscape}">')
            f.write(f'<td><b>[UNMATCHED]</b><br>{orig_name} {badges_html}</td>')
            f.write(f'<td><img src="thumbnails/local/{esc(thumb_name)}"></td>')
            f.write(f'<td>{img_path}</td>')
            f.write(f'<td>{timestamp}</td>')
            f.write('<td><button class="btn btn-warning btn-toggle-nonlive" onclick="toggleNonLive(this)">🚫 Mark Not Live</button></td>')
            f.write('</tr>\n')

        # 2. Matched Observations
        for r in matched_results:
            local_images = r.get("local_images", [])

            local_names = [str(Path(p)) for p in local_images]
            local_names_html = "<br>".join(esc(x) for x in local_names)
            local_image_data = r.get("local_image_data", [])
            row_camera_types = sorted({image.get("camera_type") for image in local_image_data if image.get("camera_type")}, key=str.casefold)
            row_camera_names = sorted({image.get("camera_name") for image in local_image_data if image.get("camera_name")}, key=str.casefold)
            camera_types_json = esc(json.dumps(row_camera_types))
            camera_names_json = esc(json.dumps(row_camera_names))

            thumbs_html = "".join([
                f'<img src="thumbnails/local/{esc(get_thumbnail_filename(p))}">'
                for p in local_images
            ])

            inat_file = esc(r.get("inat_original_filename", ""))
            obs_id = esc(r.get("observation_id", ""))
            timestamp = esc(r.get("timestamp", ""))

            first_path = local_images[0] if local_images else ""
            dir_path = esc(str(Path(first_path).parent)) if first_path else ""

            human_flags = [img for img in local_image_data if img.get("human")]
            landscape_flags = [img for img in local_image_data if img.get("landscape")]
            flags_html = ""
            if human_flags:
                flags_html += '<div><span class="badge-human">Includes Humans</span></div>'
            if landscape_flags:
                flags_html += '<div><span class="badge-landscape">Landscape / no clear subject</span></div>'

            row_human = "true" if any(img.get("human") for img in local_image_data) else "false"
            row_landscape = "true" if any(img.get("landscape") for img in local_image_data) else "false"

            f.write(f'<tr data-matched="true" data-path="{esc(first_path)}" data-dir="{dir_path}" data-filename="{inat_file}" data-timestamp="{timestamp}" data-camera-types="{camera_types_json}" data-camera-names="{camera_names_json}" data-human="{row_human}" data-landscape="{row_landscape}">')
            f.write(f'<td><b>[MATCHED]</b><br>{inat_file}</td>')
            f.write(f'<td>{thumbs_html}</td>')
            f.write(f'<td><b>Local Files:</b><br>{local_names_html}{flags_html}</td>')
            f.write(f'<td><b>Obs ID:</b> {obs_id}<br><b>Date:</b> {timestamp}</td>')
            f.write('<td><span style="color:#28a745; font-weight:bold;">Matched to iNat</span></td>')
            f.write('</tr>\n')

        f.write("""</tbody>
</table>

<br>

<div class="pagination">
    <button class="btn btn-secondary" onclick="previousPage()">Previous</button>
    <span id="page_info"></span>
    <button class="btn btn-secondary" onclick="nextPage()">Next</button>
</div>

</body>
</html>
""")



def load_existing_observations(filename):

    if not os.path.exists(filename):
        return [], None

    with open(filename, "r", encoding="utf-8") as f:
        observations = json.load(f)

    if not observations:
        return [], None

    highest_id = max(obs["id"] for obs in observations)

    print(f"Loaded {len(observations)} observations")
    print(f"Continuing above observation ID {highest_id}")

    return observations, highest_id

# ======================================================
# Download observations
# ======================================================


def fetch_observations(
        client,
        username,
        existing_observations=None,
        id_above=None,
        limit=None):


    observations = list(existing_observations or [])

    first_request = True

    while True:


        print(
            f"Downloading iNaturalist "
        )


        request_page_size = (
            limit
            if limit and limit < PAGE_SIZE
            else PAGE_SIZE
        )

        params = {


            "user_login":
                username,


            "per_page":
                request_page_size,


            "order_by":
                "id",


            "order":
                "asc",


            # Explicit fields

            "fields":
                ",".join(
                    [

                        "id",

                        "user.login",

                        "observed_on",

                        "observed_on_string",

                        "time_observed_at",

                        "place_guess",

                        "lat",

                        "lng",

                        "geojson",

                        "photos.id",

                        "photos.original_filename",

                        "taxon.id",

                    ]
                )

        }

        if id_above:
            params["id_above"] = id_above

        data = client.get(
            "/observations",
            params
        )

        if first_request:

            print(
                "iNaturalist total_results:",
                data.get("total_results")
            )

            first_request = False

        batch = data.get(
            "results",
            []
        )


        if not batch:
            print("BREAK: empty batch")
            break

        id_above = max(
            obs["id"]
            for obs in batch
        )

        for obs in batch:


            # Safety check:
            # never accept another user's observation

            obs_user = obs.get(
                "user",
                {}
            )


            if DEBUG_MODE:

                print("\n========== DEBUG OBSERVATION ==========")

                print(
                    json.dumps(
                        obs,
                        indent=2,
                        ensure_ascii=False
                    )
                )

                print(
                    "=======================================\n"
                )


            if (
                obs_user.get("login")
                !=
                username
            ):

                raise RuntimeError(

                    "\nReceived an observation "
                    "that does not belong to the authenticated user\n\n"

                    f"Requested username: {username}\n"

                    f"Returned user id: {obs_user.get('id')}\n"

                    f"Returned login: {obs_user.get('login')}\n"

                    f"Returned name: {obs_user.get('name')}\n"

                )


            # Mandatory filename check

            for photo in obs.get(
                "photos",
                []
            ):


                if not photo.get(
                    "original_filename"
                ):

                    raise RuntimeError(

                        "\nMissing original_filename\n"
                        f"Observation ID: {obs.get('id')}\n"
                        f"Photo ID: {photo.get('id')}\n\n"
                        "Cannot continue safely.\n"
                        "The script does not guess filenames "
                        "from URLs."
                    )



        observations.extend(
            batch
        )

        with open("observations.json", "w", encoding="utf-8") as f:
            json.dump(
                observations,
                f,
                ensure_ascii=False
            )

        print(
            "Downloaded:",
            len(observations),
            "last id:",
            batch[-1]["id"]
        )

        # Test mode

        if (
            limit
            and
            len(observations) >= limit
        ):

            observations = (
                observations[:limit]
            )

            print("BREAK: reached limit")
            break



        total = data.get(
            "total_results"
        )

    print(
        "Total observations:",
        len(observations)
    )


    return observations



# ======================================================
# Extract photo records
# ======================================================


def parse_inaturalist_photos(
        observations):


    photos = []



    for obs in observations:


        observation_id = obs.get(
            "id"
        )


        taxon = (
            obs.get("taxon")
            or {}
        ).get(
            "name"
        )


        observation_time = None


        raw_time = obs.get(
            "time_observed_at"
        )


        if raw_time:


            try:

                observation_time = (
                    datetime.fromisoformat(
                        raw_time.replace(
                            "Z",
                            "+00:00"
                        )
                    )
                    .replace(
                        tzinfo=None
                    )
                )


            except Exception:

                pass



        for photo in obs.get(
            "photos",
            []
        ):



            filename = photo.get(
                "original_filename"
            )


            # Double safety check

            if not filename:

                raise RuntimeError(
                    f"Missing filename for "
                    f"observation {observation_id}"
                )



            photos.append(
                {

                    "observation_id":
                        observation_id,


                    "taxon":
                        taxon,


                    "filename":
                        normalize_filename(
                            filename
                        ),

                    "inat_original_filename":
                        filename,

                    "timestamp":
                        observation_time,


                    "photo_url":
                        photo.get(
                            "original_url"
                        )

                }
            )



    print(
        "Total iNaturalist photos:",
        len(photos)
    )


    return photos

# ======================================================
# Local image scanning
# ======================================================


def is_supported_image(
        path):


    return (

        os.path.isfile(path)

        and

        os.path.splitext(path)[1].lower()
        in SUPPORTED_EXTENSIONS

    )



def get_file_signature(
        path):


    stat = os.stat(
        path
    )


    return {

        "size":
            stat.st_size,

        "mtime":
            stat.st_mtime

    }



def extract_exif_datetime(
        path):


    try:

        img = Image.open(
            path
        )


        exif = img.getexif()


        if not exif:

            return None



        for key, value in exif.items():


            tag = ExifTags.TAGS.get(
                key
            )


            if tag in (

                "DateTimeOriginal",

                "DateTime"

            ):


                return datetime.strptime(

                    value,

                    "%Y:%m:%d %H:%M:%S"

                )



    except Exception:

        pass



    return None


def extract_exif_camera(
        path):

    try:
        img = Image.open(path)
        exif = img.getexif()
        values = {ExifTags.TAGS.get(key): value for key, value in exif.items()}

        def clean(value):
            if isinstance(value, bytes):
                value = value.decode("utf-8", errors="replace")
            return str(value).strip() if value else None

        return {
            "camera_type": clean(values.get("Make")),
            "camera_name": clean(values.get("Model"))
        }
    except Exception:
        return {"camera_type": None, "camera_name": None}


_PERSON_MODEL = None
_CLIP_MODEL = None
_CLIP_PROCESSOR = None


def _load_person_model():
    global _PERSON_MODEL
    if _PERSON_MODEL is not None:
        return _PERSON_MODEL
    if YOLO is None:
        return None
    try:
        _PERSON_MODEL = YOLO("yolov8n.pt")
    except Exception:
        return None
    return _PERSON_MODEL


def _load_clip_model():
    global _CLIP_MODEL, _CLIP_PROCESSOR
    if _CLIP_MODEL is not None and _CLIP_PROCESSOR is not None:
        return _CLIP_MODEL, _CLIP_PROCESSOR
    if AutoProcessor is None or AutoModel is None:
        return None, None
    try:
        model_name = "openai/clip-vit-large-patch14"
        processor = AutoProcessor.from_pretrained(model_name)
        model = AutoModel.from_pretrained(model_name)
        model.eval()
        _CLIP_PROCESSOR = processor
        _CLIP_MODEL = model
        return model, processor
    except Exception:
        return None, None


def detect_photo_flags(path):
    """Detect likely people and landscapes with local ML models when available.

    Uses YOLOv8 for person detection and CLIP for landscape detection.
    """
    likely_human = False
    likely_landscape = False

    try:
        with Image.open(path) as img:
            rgb = img.convert("RGB")
            w, h = rgb.size
            if w <= 0 or h <= 0:
                return {"human": False, "landscape": False}
    except Exception:
        return {"human": False, "landscape": False}

    # Detect humans with YOLOv8
    person_model = _load_person_model()
    if person_model is not None:
        try:
            results = person_model(path, verbose=False, conf=0.35)
            if results:
                result = results[0]
                names = getattr(result, "names", {}) or {}
                boxes = getattr(result, "boxes", None)
                if boxes is not None:
                    for box in boxes:
                        cls_index = int(box.cls.item()) if hasattr(box.cls, "item") else int(box.cls)
                        class_name = names.get(cls_index, "").lower()
                        confidence = float(box.conf.item()) if hasattr(box.conf, "item") else float(box.conf)
                        if class_name == "person" and confidence >= 0.35:
                            likely_human = True
                            break
        except Exception:
            likely_human = False

    # Detect landscapes with CLIP when human is not found.
    if not likely_human:
        clip_model, clip_processor = _load_clip_model()
        if clip_model is not None and clip_processor is not None:
            try:
                image = Image.open(path).convert("RGB")
                landscape_prompts = [
                    "a landscape photo",
                    "a scenic outdoor landscape",
                    "a mountain view",
                    "a wide open nature scene",
                    "a forest or countryside view",
                ]
                negative_prompts = [
                    "a photo of a person",
                    "a person in a group",
                    "an indoor room",
                    "a food court or indoor restaurant",
                    "a city street",
                    "a building interior",
                ]
                texts = landscape_prompts + negative_prompts
                inputs = clip_processor(
                    text=texts,
                    images=image,
                    return_tensors="pt",
                    padding=True,
                )
                with torch.no_grad():
                    outputs = clip_model(**inputs)
                    logits = outputs.logits_per_image
                    scores = logits[0].cpu().numpy()

                positive = max(float(score) for score in scores[:len(landscape_prompts)])
                negative = max(float(score) for score in scores[len(landscape_prompts):])

                # Calibrated against the labeled examples in examples_of_landscapes.txt.
                # For the ViT-L/14 model, threshold 2.0 produced the best separation on that curated set.
                likely_landscape = (positive - negative) > 2.0
            except Exception:
                likely_landscape = False

    return {
        "human": bool(likely_human),
        "landscape": bool(likely_landscape),
    }


def scan_local_images(
        folders,
        use_cache=False,
        non_live_file="non_live_photos.json"):

    def make_sig_key(sig):
        if isinstance(sig, dict):
            return (sig.get("size"), sig.get("mtime"))
        if isinstance(sig, (list, tuple)) and len(sig) == 2:
            return (sig[0], sig[1])
        return sig

    non_live_paths = set()
    non_live_signatures = set()

    if non_live_file and os.path.exists(non_live_file):
        try:
            nl_data = load_json(non_live_file)
            if isinstance(nl_data, dict):
                non_live_paths = {os.path.abspath(p) for p in nl_data.get("paths", [])}
                non_live_signatures = {make_sig_key(s) for s in nl_data.get("signatures", []) if s is not None}
            elif isinstance(nl_data, list):
                non_live_paths = {os.path.abspath(p) for p in nl_data}
            print(f"Loaded non-live creature exclusions from {non_live_file}.")
        except Exception as e:
            print(f"Warning: Failed to load {non_live_file}: {e}")

    cache = {}
    if use_cache and os.path.exists(LOCAL_CACHE_FILE):
        cache = load_json(LOCAL_CACHE_FILE)

    images = []
    updated = False
    skipped_non_live = 0
    image_files = []
    for folder in folders:
        for root, _, files in os.walk(folder):
            for filename in files:
                path = os.path.join(root, filename)
                if is_supported_image(path):
                    image_files.append(path)

    for path in tqdm(image_files, desc="Scanning local images", unit="image"):
        filename = os.path.basename(path)
        abs_path = os.path.abspath(path)
        signature = get_file_signature(path)
        sig_key = (signature.get("size"), signature.get("mtime"))
        norm_name = normalize_filename(filename)

        if (abs_path in non_live_paths or
            path in non_live_paths or
            sig_key in non_live_signatures):
            skipped_non_live += 1
            continue

        if use_cache:
            cached = cache.get(path)
            if cached and cached.get("signature") == signature:
                data = cached["data"]
                if "camera_type" not in data or "camera_name" not in data:
                    data.update(extract_exif_camera(path))
                    updated = True
                if data.get("filename") != norm_name:
                    data["filename"] = norm_name
                    updated = True
                images.append(data)
                continue

        dt = extract_exif_datetime(path)
        camera = extract_exif_camera(path)
        flags = {"human": False, "landscape": False}
        data = {
            "path": path,
            "filename": norm_name,
            "original_filename": filename,
            "timestamp": dt.isoformat() if dt else None,
            "timestamp_minute": timestamp_minute(dt),
            "human": flags.get("human", False),
            "landscape": flags.get("landscape", False)
        }
        data.update(camera)

        if use_cache:
            cache[path] = {
                "signature": signature,
                "data": data
            }
            updated = True

        images.append(data)

    if use_cache and updated:
        save_json(LOCAL_CACHE_FILE, cache)

    if skipped_non_live > 0:
        print(f"Skipped {skipped_non_live} local photos marked as non-live creatures.")

    print("Local images indexed:", len(images))
    return images



# ======================================================
# Build filename index
# ======================================================


def build_local_indexes(
        local_images):


    filename_index = {}


    timestamp_index = {}



    for image in local_images:


        filename = image.get(
            "filename"
        )


        if filename:


            filename_index.setdefault(
                filename,
                []
            ).append(
                image
            )



        timestamp = image.get(
            "timestamp_minute"
        )


        if timestamp:


            timestamp_index.setdefault(
                timestamp,
                []
            ).append(
                image
            )



    return (
        filename_index,
        timestamp_index
    )

# ======================================================
# Matching engine
# ======================================================


def match_single_photo(
        inat_photo,
        filename_index):


    filename = inat_photo.get(
        "filename"
    )


    if not filename:


        raise RuntimeError(
            "Internal error: "
            "iNaturalist photo has no filename"
        )



    candidates = filename_index.get(
        normalize_filename(filename),
        []
    )



    # --------------------------------------------------
    # No local file with this name
    # --------------------------------------------------

    if not candidates:


        return {

            "status":
                "NOT_FOUND",

            "local_image":
                None,

            "match_method":
                None

        }



    # --------------------------------------------------
    # Exactly one filename match
    # --------------------------------------------------

    if len(candidates) == 1:


        return {

            "status":
                "MATCHED",

            "local_image":
                candidates[0],

            "match_method":
                "filename"

        }



    # --------------------------------------------------
    # Multiple identical filenames
    # Use timestamp
    # --------------------------------------------------


    inat_dt = inat_photo.get("timestamp")
    inat_time = timestamp_minute(inat_dt)

    if inat_time:
        timestamp_matches = []
        for candidate in candidates:
            if candidate.get("timestamp_minute") == inat_time:
                timestamp_matches.append(candidate)

        if len(timestamp_matches) == 1:
            return {
                "status": "MATCHED",
                "local_image": timestamp_matches[0],
                "match_method": "filename+timestamp"
            }
        elif len(timestamp_matches) > 1:
            return {
                "status": "AMBIGUOUS_TIMESTAMP",
                "local_image": timestamp_matches,
                "match_method": None
            }

        # Fallback: Timezone/DST shift tolerance (within +/- 2 hours) or same date match
        if inat_dt:
            near_matches = []
            for candidate in candidates:
                cand_ts_str = candidate.get("timestamp")
                if cand_ts_str:
                    try:
                        cand_dt = datetime.fromisoformat(cand_ts_str)
                        diff = abs((cand_dt - inat_dt).total_seconds())
                        if diff <= 7200:  # within 2 hours (1-2 hr timezone/DST shift)
                            near_matches.append(candidate)
                    except Exception:
                        pass

            if len(near_matches) == 1:
                return {
                    "status": "MATCHED",
                    "local_image": near_matches[0],
                    "match_method": "filename+near_timestamp"
                }

            # Fallback: Same date match (using EXIF timestamp or folder path date)
            inat_date_str = inat_dt.strftime("%Y-%m-%d")
            date_matches = [
                c for c in candidates
                if get_candidate_date(c) == inat_date_str
            ]
            if len(date_matches) == 1:
                return {
                    "status": "MATCHED",
                    "local_image": date_matches[0],
                    "match_method": "filename+same_date"
                }

            # Fallback: Closest timestamp match if best candidate is within 48h and significantly closer than others
            cand_diffs = []
            for candidate in candidates:
                cand_ts_str = candidate.get("timestamp")
                if cand_ts_str:
                    try:
                        cand_dt = datetime.fromisoformat(cand_ts_str)
                        diff_seconds = abs((cand_dt - inat_dt).total_seconds())
                        cand_diffs.append((diff_seconds, candidate))
                    except Exception:
                        pass

            if cand_diffs:
                cand_diffs.sort(key=lambda x: x[0])
                best_diff, best_cand = cand_diffs[0]
                second_best_diff = cand_diffs[1][0] if len(cand_diffs) > 1 else float("inf")
                # If best match is within 48h and at least 12h closer than the next best candidate
                if best_diff <= 172800 and (second_best_diff - best_diff >= 43200):
                    return {
                        "status": "MATCHED",
                        "local_image": best_cand,
                        "match_method": "filename+closest_timestamp"
                    }



    # --------------------------------------------------
    # Filename duplicated but timestamp
    # did not resolve it
    # --------------------------------------------------

    return {


        "status":
            "AMBIGUOUS_FILENAME",


        "local_image":
            candidates,


        "match_method":
            None

    }



# ======================================================
# Match all iNaturalist photos
# ======================================================


def compare_photos(
        inat_photos,
        filename_index):


    results = []

    for photo in tqdm(inat_photos, desc="Matching iNat photos", unit="photo"):
        result = match_single_photo(
            photo,
            filename_index
        )



        local = result.get(
            "local_image"
        )



        # Convert image object(s)
        # to output-friendly values


        if isinstance(
            local,
            list
        ):


            local_paths = [

                x["path"]

                for x in local

            ]

        elif local:


            local_paths = [
                local["path"]
            ]

        else:

            local_paths = []

        local_image_data = local if isinstance(local, list) else ([local] if local else [])



        results.append(
            {


                "observation_id":
                    photo["observation_id"],


                "taxon":
                    photo["taxon"],


                "inat_filename":
                    photo["filename"],

                "inat_original_filename":
                    photo["inat_original_filename"],

                "inat_timestamp":
                    timestamp_minute(
                        photo["timestamp"]
                    ),


                "inat_photo_url":
                    photo["photo_url"],


                "local_images":
                    local_paths,

                "local_image_data":
                    local_image_data,


                "status":
                    result["status"],


                "match_method":
                    result["match_method"]

            }
        )



    return results

# ======================================================
# XLSX generation
# ======================================================


def create_thumbnail(
        image_path,
        row_number):


    try:


        img = Image.open(
            image_path
        )
        img = ImageOps.exif_transpose(img)


        img.thumbnail(
            THUMB_SIZE
        )


        thumb_path = (
            f"/tmp/"
            f"inat_thumbnail_{row_number}.jpg"
        )

        if os.path.exists(thumb_path) and os.path.getsize(thumb_path) > 0:
            return thumb_path


        img.convert(
            "RGB"
        ).save(
            thumb_path,
            quality=90
        )


        return thumb_path



    except Exception as e:


        print(
            f"Cannot create thumbnail "
            f"for {image_path}: {e}"
        )


        return None


# ======================================================
# XLSX generation (two sheets)
# ======================================================


def create_thumbnail(
        image_path,
        row_number):

    try:

        img = Image.open(
            image_path
        )

        img.thumbnail(
            THUMB_SIZE
        )


        thumb_path = (
            f"/tmp/"
            f"inat_thumbnail_{row_number}.jpg"
        )


        img.convert(
            "RGB"
        ).save(
            thumb_path,
            quality=90
        )


        return thumb_path


    except Exception:

        return None



def create_xlsx(
        results,
        local_missing,
        output_file):


    wb = Workbook()



    # ==================================================
    # Sheet 1
    # ==================================================

    ws = wb.active

    ws.title = (
        "iNaturalist comparison"
    )


    headers = [

        "Observation ID",

        "Taxon",

        "iNaturalist Filename",

        "iNaturalist Timestamp",

        "Local Image(s)",

        "Status",

        "Match Method",

        "Thumbnail"

    ]


    ws.append(
        headers
    )


    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )


    row = 2


    for result in results:


        local_images = result[
            "local_images"
        ]


        ws.append(
            [

                result["observation_id"],

                result["taxon"],

                result["inat_original_filename"],

                result["inat_timestamp"],

                "\n".join(
                    local_images
                ),

                result["status"],

                result["match_method"]
                or "",

                ""

            ]
        )



        if (

            result["status"]
            ==
            "MATCHED"

            and

            len(local_images)
            ==
            1

        ):


            thumb = create_thumbnail(

                local_images[0],

                row

            )


            if thumb:


                image = XLImage(
                    thumb
                )


                image.width = 120

                image.height = 120


                ws.add_image(

                    image,

                    f"H{row}"

                )


                ws.row_dimensions[
                    row
                ].height = 95



        row += 1



    for col, width in {

        "A":16,
        "B":35,
        "C":35,
        "D":22,
        "E":70,
        "F":22,
        "G":22,
        "H":18

    }.items():

        ws.column_dimensions[
            col
        ].width = width



    ws.freeze_panes = "A2"



    # ==================================================
    # Sheet 2
    # ==================================================

    ws2 = wb.create_sheet(
        "Local not in iNaturalist"
    )


    ws2.append(
        [
            "Filename",
            "Path",
            "Timestamp",
            "Thumbnail"
        ]
    )


    row = 2


    for image in local_missing:

        ws2.append(
            [
                image["original_filename"],
                image["path"],
                image["timestamp_minute"],
                ""
            ]
        )


        thumb = create_thumbnail(
            image["path"],
            row
        )


        if thumb:

            xl_image = XLImage(
                thumb
            )

            xl_image.width = 120
            xl_image.height = 120


            ws2.add_image(
                xl_image,
                f"D{row}"
            )


            ws2.row_dimensions[
                row
            ].height = 95


        row += 1

    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 90
    ws2.column_dimensions["C"].width = 22
    ws2.column_dimensions["D"].width = 18

    wb.save(
        output_file
    )


    print(
        "Created:",
        output_file
    )


# ======================================================
# Convert XLSX to Apple Numbers
# ======================================================


def convert_to_numbers(
        xlsx_file):


    if not sys.platform.startswith(
        "darwin"
    ):

        print(
            "Numbers conversion "
            "requires macOS"
        )

        return



    numbers_file = (
        os.path.splitext(
            xlsx_file
        )[0]
        +
        ".numbers"
    )



    applescript = f'''

tell application "Numbers"

    activate

    open POSIX file "{os.path.abspath(xlsx_file)}"

    delay 5

    save front document as POSIX file "{os.path.abspath(numbers_file)}"

    close front document

end tell

'''



    subprocess.run(
        [
            "osascript",
            "-e",
            applescript
        ]
    )



    print(
        "Created Numbers file:",
        numbers_file
    )



# ======================================================
# Add local images not found in iNaturalist
# ======================================================


def _normalized_image_path(path):
    if not path:
        return ""
    try:
        return os.path.abspath(os.path.normpath(path))
    except Exception:
        return str(path)


def apply_heuristic_flags_to_unmatched(
        local_images,
        results,
    detection_cache_file=PHOTO_DETECTION_CACHE_FILE,
    detect_missing=True):
    matched_paths = set()
    for result in results:
        if result.get("status") == "MATCHED":
            for path in result.get("local_images", []):
                matched_paths.add(_normalized_image_path(path))

    unmatched_images = [
        image
        for image in local_images
        if _normalized_image_path(image.get("path")) not in matched_paths
    ]

    detection_cache = {}
    if detection_cache_file and os.path.exists(detection_cache_file):
        try:
            detection_cache = load_json(detection_cache_file)
            if not isinstance(detection_cache, dict):
                detection_cache = {}
            print("Loaded photo detection results from", detection_cache_file)
        except Exception as e:
            print(f"Warning: Failed to load {detection_cache_file}: {e}")

    updated = False
    cached_count = 0
    detected_count = 0
    action = "detection" if detect_missing else "cache loading"
    print(
        f"Running human/landscape {action} on unmatched local images: "
        f"{len(unmatched_images)} of {len(local_images)}"
    )

    for image in tqdm(
            unmatched_images,
            desc=(
                "Trying to identify humans and landscapes"
                if detect_missing
                else "Loading cached human/landscape results"
            ),
            unit="image"):
        path = image.get("path")
        cache_key = _normalized_image_path(path)
        signature = get_file_signature(path)
        cached = detection_cache.get(cache_key)
        if cached and cached.get("signature") == signature:
            flags = cached.get("flags", {})
            cached_count += 1
        elif detect_missing:
            flags = detect_photo_flags(path)
            detection_cache[cache_key] = {
                "signature": signature,
                "flags": {
                    "human": bool(flags.get("human", False)),
                    "landscape": bool(flags.get("landscape", False)),
                },
            }
            updated = True
            detected_count += 1
        else:
            flags = {"human": False, "landscape": False}
        image["human"] = flags.get("human", False)
        image["landscape"] = flags.get("landscape", False)

    if updated and detection_cache_file:
        save_json(detection_cache_file, detection_cache)
        print("Saved photo detection results to", detection_cache_file)
    if cached_count:
        print("Reused cached photo detection results:", cached_count)
    if detected_count:
        print("New photo detections calculated:", detected_count)

    return local_images


def find_local_not_in_inaturalist(
        local_images,
        results):


    matched_paths = set()



    for result in results:


        if (
            result["status"]
            ==
            "MATCHED"
        ):


            for path in result["local_images"]:

                matched_paths.add(
                    _normalized_image_path(path)
                )



    missing = []



    for image in local_images:

        image_path = _normalized_image_path(image.get("path"))


        if (
            image_path
            not in
            matched_paths
        ):


            missing.append(
                image
            )



    return missing


def load_observation_ids(filename):

    ids = set()

    with open(
        filename,
        "r",
        encoding="utf-8"
    ) as f:

        for line in f:

            line = line.strip()

            if not line:
                continue

            try:

                ids.add(
                    int(line)
                )

            except ValueError:

                # Skip header or non-ID lines
                continue


    return ids

# ======================================================
# Main
# ======================================================


def main():


    parser = argparse.ArgumentParser(

        description=
        "Compare iNaturalist photos "
        "with local images"

    )



    parser.add_argument(

        "--folders",

        nargs="+",

        required=False,

        help=
        "Folders containing local photos"

    )


    parser.add_argument(
        "--username",
        required=False,
        help="Your iNaturalist username"
    )



    parser.add_argument(

        "--limit",

        type=int,

        default=None,

        help=
        "Download only N observations "
        "for testing"

    )


    parser.add_argument(
        "--output-dir",
        default="inat_comparison"
    )

    parser.add_argument(

        "--numbers",

        action="store_true",

        help=
        "Convert XLSX to Apple Numbers"

    )


    parser.add_argument(
        "--id-export",
        help="iNaturalist exported observation IDs file"
    )

    parser.add_argument(
        "--use-cache",
        action="store_true",
        help="Enable local photos caching (disabled by default)"
    )

    parser.add_argument(
        "--detect-photo-flags",
        action="store_true",
        help="Detect humans and landscapes in unmatched photos and reuse saved results"
    )

    parser.add_argument(
        "--detection-cache",
        default=None,
        help="Load saved human/landscape results from this JSON file without running detection"
    )

    parser.add_argument(
        "--import-detection-cache-from-report",
        metavar="REPORT_HTML",
        help="Create the detection cache from an existing report.html and exit"
    )

    parser.add_argument(
        "--non-live-file",
        default="non_live_photos.json",
        help="JSON file containing list of non-live creature photos to skip scanning"
    )

    parser.add_argument(
        "--sort-by",
        choices=["date", "filename", "dir"],
        default="date",
        help="Default sorting criterion for HTML report (date, filename, dir)"
    )

    parser.add_argument(
        "--sort-order",
        choices=["asc", "desc"],
        default="asc",
        help="Default sorting order for HTML report (asc, desc)"
    )

    parser.add_argument(
        "--no-group-by-dir",
        action="store_true",
        help="Disable grouping by directory in default HTML report sorting"
    )


    args = parser.parse_args()

    if args.import_detection_cache_from_report:
        import_detection_cache_from_report(
            args.import_detection_cache_from_report,
            args.detection_cache or PHOTO_DETECTION_CACHE_FILE
        )
        return

    if not args.folders or not args.username:
        parser.error("--folders and --username are required unless importing a detection cache from report.html")

    global DEBUG_MODE

    if args.limit == 1:
        DEBUG_MODE = True



    # ------------------------------
    # Authentication
    # ------------------------------

    token = load_api_token()

    client = INaturalistClient(
        token
    )

    username = args.username


    # ------------------------------
    # Download observations
    # ------------------------------

    observations, id_above = load_existing_observations(
        "observations.json"
    )

    observations = fetch_observations(
        client,
        username,
        observations,
        id_above,
        args.limit,
    )

    if args.id_export:

        exported_ids = load_observation_ids(
            args.id_export
        )

        api_ids = {
            obs["id"]
            for obs in observations
        }


        missing = exported_ids - api_ids

        extra = api_ids - exported_ids


        print()
        print("========== ID comparison ==========")

        print(
            "Export IDs:",
            len(exported_ids)
        )

        print(
            "API IDs:",
            len(api_ids)
        )

        print(
            "Missing from API:",
            len(missing)
        )

        print(
            "Extra from API:",
            len(extra)
        )


        if missing:

            print(
                "First missing IDs:"
            )

            print(
                list(missing)[:20]
            )


        print(
            "=================================="
        )



    inat_photos = parse_inaturalist_photos(

        observations

    )



    # ------------------------------
    # Local files
    # ------------------------------


    local_images = scan_local_images(
        args.folders,
        use_cache=args.use_cache,
        non_live_file=args.non_live_file
    )



    filename_index, timestamp_index = build_local_indexes(

        local_images

    )



    # ------------------------------
    # Matching
    # ------------------------------


    results = compare_photos(

        inat_photos,

        filename_index

    )

    if args.detect_photo_flags or args.detection_cache:
        local_images = apply_heuristic_flags_to_unmatched(
            local_images,
            results,
            detection_cache_file=args.detection_cache or PHOTO_DETECTION_CACHE_FILE,
            detect_missing=args.detect_photo_flags
        )
    else:
        print(
            "Photo human/landscape detection disabled; use --detect-photo-flags "
            "to calculate flags or --detection-cache FILE to load saved flags."
        )



    matched = sum(

        1

        for r in results

        if r["status"]
        ==
        "MATCHED"

    )



    print()

    print(
        "Results:"
    )

    print(
        "Total iNaturalist photos:",
        len(results)
    )

    print(
        "Matched:",
        matched
    )



    # ------------------------------
    # Output
    # ------------------------------


    local_missing = find_local_not_in_inaturalist(
        local_images,
        results
    )


    print(
        "Local files not in iNaturalist:",
        len(local_missing)
    )


    create_html_report(
        results,
        local_missing,
        args.output_dir,
        sort_by=args.sort_by,
        sort_order=args.sort_order,
        group_by_dir=not args.no_group_by_dir
    )


    if args.numbers:


        convert_to_numbers(

            args.output

        )




if __name__ == "__main__":

    main()
